# -*- coding: utf-8 -*-
"""製作指図書Excelから進捗テーブル用SQLを作る。"""
import datetime
import json
import re
import unicodedata
from pathlib import Path

import openpyxl

SRC = Path(r"c:\Users\S002\Downloads\D令7～8製作指図書一覧 (1).xlsx")
OUT = Path(r"c:\社内資料\to\factory-materials\migrate-work-instruction-progress.sql")

SHEETS = [
    ("製作指図書一覧", "製作"),
    ("切替依頼書一覧", "切替"),
    ("修理依頼書一覧", "修理"),
]

SHOPS = [
    ("製管", "製管"),
    ("K-1", "板切"),
    ("K-2", "溶接"),
    ("K-3", "機械"),
    ("A-1", "塗装"),
    ("A-2", "検査"),
    ("A-3", "組立"),
    ("P-1", "パネル"),
]

CATEGORY = {
    "FF0000FF": "たばこ",
    "FFFF00FF": "食品",
    "FFFF9900": "暖房機",
    "FFFF0000": "作業機",
    "FF4A86E8": "青",
}


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        return value.date().isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value).strip()


def rgb_of(cell):
    fill = cell.fill
    if not fill or not fill.patternType or fill.patternType == "none":
        return None
    fg = fill.fgColor
    if fg is None or fg.type != "rgb":
        return None
    try:
        raw = str(fg.rgb or "")
    except Exception:
        return None
    if not raw or raw == "00000000":
        return None
    return raw.upper()


def sql_text(value):
    if value is None:
        return "NULL"
    text = str(value).replace("\x00", "")
    return "'" + text.replace("'", "''") + "'"


def sql_num(value):
    if value is None or value == "":
        return "NULL"
    try:
        number = float(value)
    except (TypeError, ValueError):
        return "NULL"
    if number != number:
        return "NULL"
    return str(int(number)) if number.is_integer() else str(round(number, 4))


def sql_date(value):
    text = cell_text(value)
    if re.match(r"^\d{4}-\d{2}-\d{2}$", text):
        return sql_text(text)
    return "NULL"


def fiscal_year_of(series):
    match = re.search(r"([0-9０-９]+)\s*$", series or "")
    if not match:
        return None
    number = int(unicodedata.normalize("NFKC", match.group(1)))
    if 1 <= number <= 30:
        return 2018 + number
    return None


def plain_no(value):
    text = cell_text(value)
    if re.fullmatch(r"\d+\.0", text):
        return str(int(float(text)))
    return text


def is_section(order, name):
    if not order:
        return False
    if len(order) <= 8 and (name.startswith("令和") or "年度" in name):
        return True
    return False


def load_sheet(wb_values, wb_styles, sheet_name, doc_type):
    values = wb_values[sheet_name]
    styles = wb_styles[sheet_name]
    value_rows = list(values.iter_rows(values_only=True))
    header_idx = None
    for index, row in enumerate(value_rows[:12]):
        labels = [cell_text(item) for item in row]
        if "指図書№" in labels and "名称" in labels:
            header_idx = index
            header = labels
            break
    if header_idx is None:
        raise RuntimeError(f"header not found: {sheet_name}")

    def col(label):
        return header.index(label) if label in header else None

    order_i = col("指図書№")
    name_i = col("名称")
    model_i = col("型式")
    dept_i = col("部門")
    person_i = col("担当者")
    place_i = col("納入場所")
    due_i = col("製作期限")
    wish_i = col("完了希望")
    plan_i = col("予定台数")
    partial_i = col("分納台数")
    start_i = col("発生日時")
    elapsed_i = col("発生経過月数")
    done_i = col("完了日時")
    serial_i = col("製造番号")
    slip_i = col("入庫伝票")
    comment_i = col("コメント欄")
    received_i = col("受注№")
    shop_index = {code: col(code) for code, _name in SHOPS}

    records = []
    parent_sort = None
    parent_series = ""
    parent_fiscal = None
    parent_category = None
    sort_no = 0
    branch_i = order_i + 1 if order_i is not None else 2
    for offset, value_row in enumerate(value_rows[header_idx + 1 :], start=header_idx + 2):
        style_row = styles[offset]
        padded = list(value_row) + [None] * 8

        def at(index):
            if index is None or index >= len(padded):
                return None
            return padded[index]

        order = cell_text(at(order_i))
        name = cell_text(at(name_i))
        branch = cell_text(at(branch_i))
        if order in {"暖房機", "作業機", "たばこ", "食品", "指図書№"}:
            continue
        if name in {"名称"}:
            continue
        if is_section(order, name):
            parent_sort = None
            parent_category = None
            continue
        if not order and not name and not branch:
            continue

        category = CATEGORY.get(rgb_of(style_row[0]) or "")
        shops = []
        for code, shop_name in SHOPS:
            index = shop_index.get(code)
            raw = at(index) if index is not None else None
            completed = raw is True or cell_text(raw).upper() == "TRUE"
            fill = None
            if index is not None and index < len(style_row):
                fill = rgb_of(style_row[index])
            assigned = fill == "FFFFFF00" or completed
            shops.append(
                {
                    "code": code,
                    "name": shop_name,
                    "assigned": assigned,
                    "completed": completed,
                }
            )

        if order:
            sort_no += 1
            parent_sort = sort_no
            parent_series = order
            parent_fiscal = fiscal_year_of(order)
            parent_category = category or None
            records.append(
                build_record(
                    doc_type, sort_no, None, order, parent_fiscal, name, "", plain_no(branch), category, shops, at,
                    model_i, dept_i, person_i, place_i, due_i, wish_i, plan_i, partial_i,
                    start_i, elapsed_i, done_i, serial_i, slip_i, comment_i, received_i,
                )
            )
        else:
            if parent_sort is None:
                continue
            sort_no += 1
            records.append(
                build_record(
                    doc_type, sort_no, parent_sort, parent_series, parent_fiscal, name, branch, "",
                    category or parent_category, shops, at,
                    model_i, dept_i, person_i, place_i, due_i, wish_i, plan_i, partial_i,
                    start_i, elapsed_i, done_i, serial_i, slip_i, comment_i, received_i,
                )
            )
    return records


def build_record(
    doc_type, sort_no, parent_sort, order, fiscal_year, name, branch, item_no, category, shops, at,
    model_i, dept_i, person_i, place_i, due_i, wish_i, plan_i, partial_i,
    start_i, elapsed_i, done_i, serial_i, slip_i, comment_i, received_i,
):
    slip_raw = at(slip_i)
    receipt = slip_raw is True or cell_text(slip_raw).upper() == "TRUE"
    elapsed_raw = at(elapsed_i)
    elapsed = None
    if isinstance(elapsed_raw, (int, float)) and not isinstance(elapsed_raw, bool):
        elapsed = int(elapsed_raw)
    due_raw = at(due_i)
    due_text = "" if isinstance(due_raw, (datetime.date, datetime.datetime)) else cell_text(due_raw)
    wish_raw = at(wish_i)
    wish_text = "" if isinstance(wish_raw, (datetime.date, datetime.datetime)) else cell_text(wish_raw)
    comment = cell_text(at(comment_i))
    if len(comment) > 1000:
        comment = comment[:1000]
    return {
        "doc_type": doc_type,
        "series_no": order,
        "fiscal_year": fiscal_year,
        "sort_no": sort_no,
        "parent_sort_no": parent_sort,
        "branch_no": branch,
        "item_no": item_no,
        "category": category or "",
        "product_name": name,
        "model": cell_text(at(model_i)),
        "department": cell_text(at(dept_i)),
        "assignee": cell_text(at(person_i)),
        "delivery_place": cell_text(at(place_i)),
        "due_on": cell_text(due_raw) if isinstance(due_raw, (datetime.date, datetime.datetime)) else "",
        "due_text": due_text,
        "wish_text": wish_text,
        "planned_qty": at(plan_i) if isinstance(at(plan_i), (int, float)) and not isinstance(at(plan_i), bool) else None,
        "partial_qty": at(partial_i) if isinstance(at(partial_i), (int, float)) and not isinstance(at(partial_i), bool) else None,
        "occurred_on": cell_text(at(start_i)) if isinstance(at(start_i), (datetime.date, datetime.datetime)) else "",
        "elapsed_months": elapsed,
        "completed_on": cell_text(at(done_i)) if isinstance(at(done_i), (datetime.date, datetime.datetime)) else "",
        "serial_no": cell_text(at(serial_i)),
        "receipt_posted": receipt,
        "comment": comment,
        "received_order_no": cell_text(at(received_i)).replace(".0", "") if cell_text(at(received_i)).endswith(".0") else cell_text(at(received_i)),
        "shops": shops,
    }


def insert_sql(row):
    shops = json.dumps(row["shops"], ensure_ascii=False)
    fiscal = "NULL" if row["fiscal_year"] is None else str(row["fiscal_year"])
    parent = "NULL" if row["parent_sort_no"] is None else str(row["parent_sort_no"])
    elapsed = "NULL" if row["elapsed_months"] is None else str(row["elapsed_months"])
    return (
        "INSERT INTO work_instruction_progress ("
        "doc_type, series_no, fiscal_year, sort_no, parent_sort_no, branch_no, category, "
        "product_name, model, department, assignee, delivery_place, due_on, due_text, wish_text, "
        "planned_qty, partial_qty, occurred_on, elapsed_months, completed_on, serial_no, "
        "receipt_posted, comment, received_order_no, shops, source"
        ") VALUES ("
        f"{sql_text(row['doc_type'])}, {sql_text(row['series_no'])}, {fiscal}, {row['sort_no']}, {parent}, "
        f"{sql_text(row['branch_no'] or None)}, {sql_text(row['category'] or None)}, "
        f"{sql_text(row['product_name'] or None)}, {sql_text(row['model'] or None)}, "
        f"{sql_text(row['department'] or None)}, {sql_text(row['assignee'] or None)}, "
        f"{sql_text(row['delivery_place'] or None)}, {sql_date(row['due_on'] or None)}, "
        f"{sql_text(row['due_text'] or None)}, {sql_text(row['wish_text'] or None)}, "
        f"{sql_num(row['planned_qty'])}, {sql_num(row['partial_qty'])}, {sql_date(row['occurred_on'] or None)}, "
        f"{elapsed}, {sql_date(row['completed_on'] or None)}, {sql_text(row['serial_no'] or None)}, "
        f"{'true' if row['receipt_posted'] else 'false'}, {sql_text(row['comment'] or None)}, "
        f"{sql_text(row['received_order_no'] or None)}, {sql_text(shops)}::jsonb, 'excel'"
        ") ON CONFLICT (doc_type, sort_no) WHERE source = 'excel' DO UPDATE SET "
        "series_no = EXCLUDED.series_no, fiscal_year = EXCLUDED.fiscal_year, "
        "parent_sort_no = EXCLUDED.parent_sort_no, branch_no = EXCLUDED.branch_no, "
        "category = EXCLUDED.category, product_name = EXCLUDED.product_name, model = EXCLUDED.model, "
        "department = EXCLUDED.department, assignee = EXCLUDED.assignee, delivery_place = EXCLUDED.delivery_place, "
        "due_on = EXCLUDED.due_on, due_text = EXCLUDED.due_text, wish_text = EXCLUDED.wish_text, "
        "planned_qty = EXCLUDED.planned_qty, partial_qty = EXCLUDED.partial_qty, "
        "occurred_on = EXCLUDED.occurred_on, elapsed_months = EXCLUDED.elapsed_months, "
        "completed_on = EXCLUDED.completed_on, serial_no = EXCLUDED.serial_no, "
        "receipt_posted = EXCLUDED.receipt_posted, comment = EXCLUDED.comment, "
        "received_order_no = EXCLUDED.received_order_no, shops = EXCLUDED.shops, updated_at = now();"
    )


def main():
    wb_values = openpyxl.load_workbook(SRC, data_only=True, read_only=True)
    wb_styles = openpyxl.load_workbook(SRC, data_only=False, read_only=False)
    all_rows = []
    for sheet_name, doc_type in SHEETS:
        loaded = load_sheet(wb_values, wb_styles, sheet_name, doc_type)
        print(doc_type, len(loaded))
        all_rows.extend(loaded)

    statements = [insert_sql(row) for row in all_rows]
    chunks = []
    current = []
    size = 0
    for statement in statements:
        encoded = len(statement.encode("utf-8")) + 1
        # SQL Editor は約 0.98MB を超えると拒否する
        if current and size + encoded > 850_000:
            chunks.append(current)
            current = []
            size = 0
        current.append(statement)
        size += encoded
    if current:
        chunks.append(current)

    header = """-- 製作指図書の進捗管理
-- Supabase の SQL Editor で、まずこのファイルを実行してください。
-- そのあと data ファイルを番号順に実行すると、Excel の一覧が入ります。
-- 再実行しても Excel 取込行（source = excel）は上書きされ、画面から登録した D指令行は残ります。
--
-- 作業班の shops:
--   assigned = true  … 製造部長が割り当てた班（Excel の黄色）
--   completed = true … その班の作業が完了（Excel のチェック）

CREATE TABLE IF NOT EXISTS work_instruction_progress (
  id uuid PRIMARY KEY DEFAULT gen_random_uuid(),
  work_order_id uuid REFERENCES work_orders(id) ON DELETE SET NULL,
  doc_type text NOT NULL DEFAULT '製作',
  series_no text NOT NULL DEFAULT '',
  fiscal_year integer,
  sort_no integer,
  parent_sort_no integer,
  branch_no text,
  category text,
  product_name text,
  model text,
  department text,
  assignee text,
  delivery_place text,
  due_on date,
  due_text text,
  wish_text text,
  planned_qty numeric(18,4),
  partial_qty numeric(18,4),
  occurred_on date,
  elapsed_months integer,
  completed_on date,
  serial_no text,
  receipt_posted boolean NOT NULL DEFAULT false,
  comment text,
  received_order_no text,
  shops jsonb NOT NULL DEFAULT '[]'::jsonb,
  source text NOT NULL DEFAULT 'manual',
  created_at timestamptz NOT NULL DEFAULT now(),
  updated_at timestamptz NOT NULL DEFAULT now()
);

CREATE UNIQUE INDEX IF NOT EXISTS uq_work_instruction_progress_excel
  ON work_instruction_progress (doc_type, sort_no)
  WHERE source = 'excel';

CREATE UNIQUE INDEX IF NOT EXISTS uq_work_instruction_progress_work_order
  ON work_instruction_progress (work_order_id)
  WHERE work_order_id IS NOT NULL;

CREATE INDEX IF NOT EXISTS idx_work_instruction_progress_doc_year
  ON work_instruction_progress (doc_type, fiscal_year, sort_no);

ALTER TABLE work_instruction_progress DISABLE ROW LEVEL SECURITY;

COMMENT ON TABLE work_instruction_progress IS '製作指図書の進捗。D指令登録で追加し、製造部長が作業班を割り当てる';
COMMENT ON COLUMN work_instruction_progress.shops IS '作業班配列。assigned=割当、completed=完了。codeは製管/K-1/K-2/K-3/A-1/A-2/A-3/P-1';
COMMENT ON COLUMN work_instruction_progress.series_no IS '指図書番号。Excelでは D令7 のような年度の束。システム登録のD指令は指令番号';
COMMENT ON COLUMN work_instruction_progress.partial_qty IS '分納台数。残台数は planned_qty - partial_qty';

-- このファイルは表の作成だけです。
-- 続けて次のファイルを番号順に SQL Editor で実行すると、Excel の一覧が入ります。
"""
    listing = "\n".join(
        f"--   {index + 1}. migrate-work-instruction-progress-data-{index + 1:02d}.sql"
        for index in range(len(chunks))
    )
    OUT.write_text(header + listing + "\n", encoding="utf-8")
    for stale in OUT.parent.glob("migrate-work-instruction-progress-data-*.sql"):
        stale.unlink()
    for index, chunk in enumerate(chunks, start=1):
        path = OUT.with_name(f"migrate-work-instruction-progress-data-{index:02d}.sql")
        path.write_text(
            "-- 先に migrate-work-instruction-progress.sql を実行してから、このファイルを実行してください。\n"
            + "\n".join(chunk)
            + "\n",
            encoding="utf-8",
        )
        print(path.name, len(chunk), path.stat().st_size)
    print("rows", len(all_rows), "files", len(chunks))


def write_item_no_sql():
    wb_values = openpyxl.load_workbook(SRC, data_only=True, read_only=True)
    wb_styles = openpyxl.load_workbook(SRC, data_only=False, read_only=False)
    lines = [
        "-- 指図書の横にある № を足す。例: 指図書 K令9、№ 1",
        "-- Supabase の SQL Editor で、このファイルを1回実行してください。",
        "-- すでに画面で入れた № は上書きしません。",
        "ALTER TABLE work_instruction_progress ADD COLUMN IF NOT EXISTS item_no text;",
        "COMMENT ON COLUMN work_instruction_progress.item_no IS '指図書の横の№。切替は K令9 の 1、2、3';",
    ]
    count = 0
    for sheet_name, doc_type in SHEETS:
        for row in load_sheet(wb_values, wb_styles, sheet_name, doc_type):
            if row["parent_sort_no"] is not None or not row["item_no"]:
                continue
            count += 1
            lines.append(
                "UPDATE work_instruction_progress SET item_no = "
                + sql_text(row["item_no"])
                + " WHERE source = 'excel' AND doc_type = "
                + sql_text(row["doc_type"])
                + " AND sort_no = "
                + str(row["sort_no"])
                + " AND (item_no IS NULL OR item_no = '');"
            )
    path = OUT.with_name("migrate-work-instruction-item-no.sql")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print("item_no", count, path.name, path.stat().st_size)


if __name__ == "__main__":
    import sys

    if "--item-no-only" in sys.argv:
        write_item_no_sql()
    else:
        main()
